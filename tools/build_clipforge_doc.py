from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('deliverables/ClipForge_AI_Features_and_Workflow_Guide.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)
NAVY='10243E'; CYAN='10B9D3'; GREEN='18A957'; LIGHT='F3F7FA'; INK='152236'; MUTED='52657A'

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def border(table):
    p=table._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6'); e.set(qn('w:color'),'D9E4EC'); b.append(e)
    p.append(b)
def celltext(cell,text,bold=False,color=INK,size=9):
    cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(str(text)); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def make_table(doc,heads,rows):
    t=doc.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; border(t)
    for i,h in enumerate(heads): celltext(t.rows[0].cells[i],h,True,'FFFFFF',9); shade(t.rows[0].cells[i],NAVY)
    for n,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row): celltext(cs[i],v)
        if n%2==0:
            for c in cs: shade(c,LIGHT)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
def bullet(doc,text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(3); p.add_run(text)
def step(doc,num,title,text):
    t=doc.add_table(rows=1,cols=2); border(t); a,b=t.rows[0].cells; shade(a,CYAN); shade(b,LIGHT)
    p=a.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(str(num)); r.bold=True; r.font.size=Pt(18); r.font.color.rgb=RGBColor(255,255,255)
    p=b.paragraphs[0]; r=p.add_run(title); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=b.add_paragraph(text); p.runs[0].font.size=Pt(9.5); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document(); s=doc.sections[0]; s.top_margin=Inches(.65); s.bottom_margin=Inches(.65); s.left_margin=Inches(.7); s.right_margin=Inches(.7)
st=doc.styles; st['Normal'].font.name='Aptos'; st['Normal'].font.size=Pt(10); st['Normal'].font.color.rgb=RGBColor.from_string(INK); st['Normal'].paragraph_format.space_after=Pt(6)
for n,z,c in [('Title',31,NAVY),('Heading 1',20,NAVY),('Heading 2',13,CYAN),('Heading 3',11,GREEN)]:
    st[n].font.name='Aptos'; st[n].font.size=Pt(z); st[n].font.bold=True; st[n].font.color.rgb=RGBColor.from_string(c)
head=s.header.paragraphs[0]; head.text='CLIPFORGE AI  |  FEATURES AND WORKFLOW'; head.runs[0].font.size=Pt(8); head.runs[0].font.bold=True; head.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
foot=s.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; foot.add_run('ClipForge AI Creator Studio').font.size=Pt(8)

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(55); r=p.add_run('CLIPFORGE AI'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(CYAN)
p=doc.add_paragraph(style='Title'); p.add_run('Features and Workflow Guide')
p=doc.add_paragraph('A clear guide to creating, polishing, packaging, and publishing short videos.'); p.runs[0].font.size=Pt(14); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
make_table(doc,['AT A GLANCE','WHAT YOU CAN DO'],[('Input','Upload one video, several videos, a project video, or a public link.'),('Output','Shorts, captions, thumbnails, metadata, music, reports, and ZIP package.'),('Publishing','Select shorts, choose thumbnails, use generated metadata, and publish to YouTube.')])
doc.add_paragraph('This guide is written for creators, editors, educators, agencies, and anyone who wants one simple path from a source video to ready-to-publish short content.')
doc.add_page_break()

doc.add_heading('What ClipForge AI Includes',1)
doc.add_paragraph('ClipForge AI is a Creator Studio with one connected workflow. You give it a source video, choose your style, and receive a complete content package. Each part stays connected to the same job so you can preview, check, download, or publish without moving between separate tools.')
make_table(doc,['AREA','FEATURES INCLUDED'],[('Creator Studio','Five guided steps: Input Source, Output Setup, Captions and Style, Music and Package, and Generate.'),('Clip creation','Semantic AI selection, fixed-duration clips, manual ranges, raw-footage mode, and short-input fallback.'),('Captions','English captions, Roman Urdu/Hindi captions, style presets, font, size, position, case, and burned-in output.'),('Visual package','Multiple thumbnail variations, preview cards, selected thumbnail upload, and downloadable images.'),('Metadata','Titles, hooks, descriptions, hashtags, keywords, platform, language, and clip range.'),('Audio','Music library, category selection, preview, volume control, and track fallback notice.'),('Results','Shorts, Thumbnails, Metadata, and Publish to YouTube tabs.'),('YouTube','OAuth connection, multi-select upload, thumbnail choice per short, visibility, and scheduling.')])

doc.add_heading('The Creator Studio Journey',1)
step(doc,1,'Input Source','Upload one video, select a project video, upload several videos, or paste a supported public link.')
step(doc,2,'Output Setup','Choose platform, aspect ratio, quality, segment mode, and clip length when required.')
step(doc,3,'Captions and Style','Turn captions on, choose the language behavior and preset, adjust the look, and preview it.')
step(doc,4,'Music and Package','Choose music, category, editing style, volume, and review the package summary.')
step(doc,5,'Generate','Start processing and watch the live stages from upload through the ZIP package.')

doc.add_heading('Input Source Features',1)
doc.add_heading('Upload Video',2); doc.add_paragraph('Choose one video for a focused run or several videos for batch processing. The upload area accepts common video files and keeps the selected source visible before you continue.')
doc.add_heading('Project Video Library',2); doc.add_paragraph('Refresh the project input folder and choose a video already available to the backend. This is useful when files are copied into the project before opening the browser.')
doc.add_heading('Paste Link',2); doc.add_paragraph('Paste a supported public video link. ClipForge downloads the source, processes it, and connects the results to the job ID.')
for x in ['Every run receives a job ID in the Status section.','The current job can be resumed after a browser refresh.','Batch runs keep separate output folders for each source video.']: bullet(doc,x)

doc.add_heading('Output Setup Features',1)
make_table(doc,['CONTROL','OPTIONS','WHAT IT CONTROLS'],[('Segment mode','Semantic AI','Finds meaningful moments and selects strong short segments.'),('Segment mode','Fixed Duration','Creates clips around the duration you enter.'),('Segment mode','Manual Ranges','Uses exact start and end ranges created by you.'),('Segment mode','Raw Footage','Finds useful speech-rich sections from less-edited footage.'),('Platform','YouTube Shorts, Instagram Reels, TikTok','Applies platform packaging choices.'),('Aspect ratio','9:16, 16:9, 1:1, 4:5','Sets the canvas shape.'),('Quality','1080p, 720p, 480p','Sets resolution and file-size balance.')])

doc.add_heading('Captions and Style Features',1)
doc.add_paragraph('Captions are part of the finished video. ClipForge can prepare the caption track and burn it into the exported short so the text remains visible on supported platforms.')
make_table(doc,['SETTING','WHAT YOU CAN CONTROL'],[('Caption mode','Captions on or off; automatic English and Roman behavior.'),('Language style','English captions or Roman Urdu/Hindi captions while source timing stays aligned.'),('Roman accuracy','OpenAI-assisted Romanization when available, with local fallback when the API is unavailable or quota is finished.'),('Preset','Clean White and other visual presets.'),('Font','Preset fonts and available local font families.'),('Position','Bottom center and other configured positions.'),('Size and case','Small to extra-large sizing and preset or custom case.'),('Preview','Preview text and video styling before the final run.')])
doc.add_paragraph('English words inside Roman captions are preserved instead of being unnecessarily changed. When the source transcript has no usable text for a selected range, the translated transcript can provide a caption fallback.')

doc.add_heading('Music and Package Features',1)
for x in ['Browse music by category and preview a track before processing.','Set music volume so speech remains clear.','Choose an editing style such as educational or another available style.','If a requested track is unavailable, the package reports the replacement.','Review a final summary showing input, platform, ratio, mode, quality, filters, captions, music, and outputs.']: bullet(doc,x)

doc.add_page_break(); doc.add_heading('Live Processing and Status',1)
doc.add_paragraph('The Status section shows the current stage, percentage, elapsed time, job ID, and progress events. Stage cards show what is complete, active, or failed.')
make_table(doc,['STAGE','WHAT HAPPENS'],[('Upload or link submitted','The source is accepted and the job is created.'),('Job queued','The backend places the job into processing.'),('Transcribing audio','Audio is extracted and speech is transcribed.'),('Detecting highlights','Clip moments are selected using the chosen mode.'),('Rendering shorts','Crop, reframe, filters, music, and captions are applied.'),('Creating captions','Caption files are prepared and burned in when enabled.'),('Generating thumbnails','Thumbnail variations are generated and saved.'),('Writing metadata','Titles, hooks, descriptions, hashtags, and keywords are written.'),('Building ZIP package','Final files are collected into one package.'),('Complete','Results are ready in the dashboard.')])

doc.add_heading('Results Dashboard',1)
doc.add_paragraph('The Results area is the main working space after processing. Summary cards show generated shorts, thumbnails, metadata files, and ZIP status. The left tabs keep each output type easy to find.')
make_table(doc,['TAB','WHAT YOU SEE AND DO'],[('Shorts','Play each generated video and check caption, audio, framing, and name.'),('Thumbnails','Review generated images and open or download variations.'),('Metadata','Preview structured fields, view raw text, and copy individual fields or the complete block.'),('Publish to YouTube','Select shorts, choose thumbnails, choose visibility or schedule, connect YouTube, and publish.')])

doc.add_heading('Thumbnail Workflow',1)
step(doc,1,'Review variations','Open Thumbnails and compare the generated images for each short.')
step(doc,2,'Select the best image','In Publish to YouTube, choose thumbnail 1, 2, or 3 for each selected short.')
step(doc,3,'Upload with the video','The selected image is sent after the video is created. PNG and JPEG files use the correct image type.')
step(doc,4,'Check channel eligibility','Custom thumbnails require YouTube channel verification. Complete phone verification if YouTube asks for it.')

doc.add_heading('Metadata Features',1)
make_table(doc,['FIELD','USE'],[('Title','Main YouTube title for the short.'),('Hook options','Opening lines that can be used as a stronger title or opening idea.'),('Description','Ready description with context and a call to action.'),('Hashtags','Hashtags in normal hashtag form.'),('Hashtags comma','Comma-separated keywords for posting tools.'),('Platform and language','Intended platform and English or Roman language.'),('Clip range','Original start and end time used for the short.'),('Keywords hit','Topics found in the source that shaped the metadata.')])

doc.add_heading('YouTube Publishing Workflow',1)
doc.add_paragraph('The Publish to YouTube tab supports one-short and multi-short publishing. ClipForge stays open while Google authentication and YouTube Studio open in separate tabs.')
for n,title,text in [(1,'Open Publish to YouTube','Open the tab after Metadata in Results.'),(2,'Select shorts','Tick one, several, or Select all shorts.'),(3,'Choose thumbnails','Pick thumbnail 1, 2, or 3 separately for each short.'),(4,'Choose settings','Choose Private, Unlisted, or Public. Select Upload now or Schedule.'),(5,'Connect YouTube','Google opens in a new tab. Select the correct account and allow upload access.'),(6,'Publish selected shorts','The button changes to Publish Selected Shorts. Video, thumbnail, and generated metadata are sent together.'),(7,'Open Studio content','After upload, YouTube Studio Content opens in a new tab for review.')]: step(doc,n,title,text)
for x in ['Private is recommended for the first test.','Scheduled uploads are sent as private videos with a publish time.','Generated title, description, tags, and category are used automatically.','The selected thumbnail is uploaded after the video is created.','The upload response includes the YouTube link and thumbnail status.']: bullet(doc,x)

doc.add_heading('Complete File Package',1)
make_table(doc,['FOLDER OR FILE','CONTENTS'],[('shorts','Final MP4 short videos.'),('captions','Caption files, including English or Roman variants.'),('thumbnails','Generated PNG/JPEG variations and prompt details.'),('meta','Per-short metadata text files.'),('reports','Run summary and processing details.'),('client_output.zip','One package containing final deliverables.')])
doc.add_heading('A Simple End to End Routine',1)
for x in ['Select a source video.','Choose YouTube Shorts, 9:16, and the required quality.','Select Semantic AI or Manual Ranges.','Turn captions on and choose the caption look.','Pick music and keep its volume below the speech.','Review the package summary and start Generate.','Wait until the Status stage shows Complete.','Play the result and check captions, audio, and framing.','Review Thumbnails and Metadata.','Open Publish to YouTube, choose shorts and thumbnails, and use Private for the first test.','Open YouTube Studio in the new tab and check the upload.']: bullet(doc,x)
doc.add_heading('Quick Reference',1)
make_table(doc,['TASK','WHERE TO GO'],[('Change source','Input Source'),('Change crop or quality','Output Setup'),('Change captions','Captions and Style'),('Change music','Music and Package'),('Check progress','Status'),('Play final video','Results > Shorts'),('Choose thumbnail','Results > Publish to YouTube'),('Copy posting text','Results > Metadata'),('Download everything','Download All as ZIP'),('Publish or schedule','Results > Publish to YouTube')])
p=doc.add_paragraph('ClipForge AI brings the complete short-video workflow into one clear Creator Studio: source in, polished shorts out, and publishing ready when you are.'); p.style='Intense Quote'
doc.save(OUT); print(OUT.resolve())
